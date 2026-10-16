"""Tests for the zip-bomb pre-check that runs before .xlsx/.xlsm/.ods files are parsed."""

from __future__ import annotations

import shutil
import zipfile
from pathlib import Path

import openpyxl
import pytest

import doormat


@pytest.fixture
def workbook(tmp_path: Path) -> Path:
    """A small, valid workbook with one credential."""
    wb = openpyxl.Workbook()
    wb.active["A1"] = "Password"
    wb.active["B1"] = "s3cret!!"
    path = tmp_path / "book.xlsx"
    wb.save(path)
    wb.close()
    return path


def _with_bomb(src: Path, dst: Path, size: int = 20 * 1024 * 1024) -> Path:
    """Copy ``src`` with its worksheet replaced by ``size`` bytes of repetitive data."""
    with zipfile.ZipFile(src) as zin, zipfile.ZipFile(dst, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "xl/worksheets/sheet1.xml":
                data = b"<worksheet>" + b" " * size + b"</worksheet>"
            zout.writestr(item.filename, data)
    return dst


def test_zip_bomb_is_refused(workbook: Path, tmp_path: Path) -> None:
    bomb = _with_bomb(workbook, tmp_path / "bomb.xlsx")
    assert bomb.stat().st_size < 200_000  # tiny on disk, 20 MiB unpacked
    with pytest.raises(ValueError, match="zip bombs"):
        doormat.load(bomb)


def test_zip_bomb_is_refused_even_without_size_limit(workbook: Path, tmp_path: Path) -> None:
    bomb = _with_bomb(workbook, tmp_path / "bomb.xlsx")
    with pytest.raises(ValueError, match="zip bombs"):
        doormat.load(bomb, max_uncompressed_bytes=None)


def test_zip_bomb_is_refused_for_ods(workbook: Path, tmp_path: Path) -> None:
    bomb = _with_bomb(workbook, tmp_path / "bomb.xlsx")
    ods = shutil.copy(bomb, tmp_path / "bomb.ods")
    with pytest.raises(ValueError, match="zip bombs"):
        doormat.load(ods)


def test_too_many_parts_is_refused(tmp_path: Path) -> None:
    path = tmp_path / "parts.xlsx"
    with zipfile.ZipFile(path, "w") as archive:
        for i in range(10_001):
            archive.writestr(f"part{i}.xml", b"")
    with pytest.raises(ValueError, match="archive parts"):
        doormat.load(path)


def test_size_limit_is_configurable(workbook: Path) -> None:
    with pytest.raises(ValueError, match="max_uncompressed_bytes"):
        doormat.load(workbook, max_uncompressed_bytes=100)
    assert len(doormat.load(workbook, max_uncompressed_bytes=None).credentials(120)) == 1
    assert len(doormat.load(workbook).credentials(120)) == 1


def test_non_zip_formats_are_not_checked(tmp_path: Path) -> None:
    path = tmp_path / "data.csv"
    path.write_text("Password,s3cret!!\n")
    assert len(doormat.load(path, max_uncompressed_bytes=1).credentials(120)) == 1


def test_far_formatted_cell_does_not_inflate_xlsx(tmp_path: Path) -> None:
    """One styled empty cell at XFD1048576 must not make the extractor walk 17 billion positions."""
    import time

    from openpyxl.styles import Font

    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "Password"
    ws["B1"] = "s3cret!!"
    ws["XFD1048576"].font = Font(bold=True)
    path = tmp_path / "far.xlsx"
    wb.save(path)
    wb.close()

    start = time.perf_counter()
    found = doormat.load(path).credentials(120)
    assert time.perf_counter() - start < 5
    assert [c.coordinate for c in found] == ["B1"]


def test_huge_column_repeat_does_not_inflate_ods(tmp_path: Path) -> None:
    """An empty cell repeated a billion times must not be expanded cell by cell."""
    import time

    pytest.importorskip("odf")
    from odf.opendocument import OpenDocumentSpreadsheet
    from odf.table import Table, TableCell, TableRow
    from odf.text import P

    doc = OpenDocumentSpreadsheet()
    table = Table(name="S1")
    row = TableRow()
    for text in ("Password", "s3cret!!"):
        cell = TableCell()
        cell.addElement(P(text=text))
        row.addElement(cell)
    row.addElement(TableCell(numbercolumnsrepeated=1_000_000_000))
    table.addElement(row)
    doc.spreadsheet.addElement(table)
    path = tmp_path / "repeat.ods"
    doc.save(str(path))

    start = time.perf_counter()
    found = doormat.load(path).credentials(120)
    assert time.perf_counter() - start < 5
    assert [c.coordinate for c in found] == ["B1"]


def _grid_workbook(tmp_path: Path, rows: int, cols: int, sheets: int = 1) -> Path:
    wb = openpyxl.Workbook()
    for s in range(sheets):
        ws = wb.active if s == 0 else wb.create_sheet(f"S{s}")
        for r in range(rows):
            ws.append([f"v{r}_{c}" for c in range(cols)])
    path = tmp_path / "grid.xlsx"
    wb.save(path)
    wb.close()
    return path


def test_cell_limit_refuses_the_file(tmp_path: Path) -> None:
    path = _grid_workbook(tmp_path, rows=30, cols=10, sheets=2)
    with pytest.raises(ValueError, match="max_cells"):
        doormat.load(path, max_cells=400)


def test_cell_limit_is_not_reported_as_a_sheet_error(tmp_path: Path) -> None:
    from doormat._limits import CellLimitError

    path = _grid_workbook(tmp_path, rows=30, cols=10, sheets=3)
    with pytest.raises(CellLimitError):
        doormat.load(path, max_cells=350)


def test_cell_limit_can_be_disabled(tmp_path: Path) -> None:
    path = _grid_workbook(tmp_path, rows=30, cols=10)
    assert doormat.load(path, max_cells=None).cell_count == 300
    assert doormat.load(path, max_cells=300).cell_count == 300


def test_cell_limit_stops_parsing_part_way_through_a_sheet(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    """The count runs while the sheet is parsed, not after it is in memory."""
    from doormat.extract import xlsx

    path = _grid_workbook(tmp_path, rows=200, cols=10)
    rows_seen = []
    real = xlsx.extract_single_sheet

    def spy(ws, *args):  # type: ignore[no-untyped-def]
        rows_seen.append(ws.max_row)
        return real(ws, *args)

    monkeypatch.setattr(xlsx, "extract_single_sheet", spy)
    with pytest.raises(ValueError, match="max_cells"):
        doormat.load(path, max_cells=100)
    assert rows_seen == []


def test_cell_limit_applies_to_csv(tmp_path: Path) -> None:
    path = tmp_path / "big.csv"
    path.write_text("\n".join(",".join(f"v{c}" for c in range(10)) for _ in range(30)))
    with pytest.raises(ValueError, match="max_cells"):
        doormat.load(path, max_cells=100)
    assert doormat.load(path).cell_count == 300


def test_cell_limit_applies_to_ods(tmp_path: Path) -> None:
    pytest.importorskip("odf")
    from odf.opendocument import OpenDocumentSpreadsheet
    from odf.table import Table, TableCell, TableRow
    from odf.text import P

    doc = OpenDocumentSpreadsheet()
    table = Table(name="S1")
    for r in range(30):
        row = TableRow()
        for c in range(10):
            cell = TableCell()
            cell.addElement(P(text=f"v{r}_{c}"))
            row.addElement(cell)
        table.addElement(row)
    doc.spreadsheet.addElement(table)
    path = tmp_path / "grid.ods"
    doc.save(str(path))
    with pytest.raises(ValueError, match="max_cells"):
        doormat.load(path, max_cells=100)
