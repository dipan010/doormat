"""Tests for doormat.extract: xlsx cell extraction via openpyxl."""

from __future__ import annotations

import zipfile
from pathlib import Path

import openpyxl
import pytest
from doormat.extract import (
    clean_comment,
    collect_merge_origins,
    extract_single_sheet,
    extract_workbook,
)
from openpyxl.chart import BarChart, Reference


@pytest.fixture
def tmp_xlsx(tmp_path: Path):
    """Helper that writes an openpyxl Workbook to a temp file and returns the path."""

    def _write(wb: openpyxl.Workbook) -> Path:
        p = tmp_path / "test.xlsx"
        wb.save(p)
        return p

    return _write


# ---------- clean_comment ----------


def test_clean_comment_strips():
    assert clean_comment("  hello  ") == "hello"


def test_clean_comment_removes_prefix():
    assert clean_comment("Comment:\nhello world") == "hello world"


def test_clean_comment_no_prefix():
    assert clean_comment("just a note") == "just a note"


def test_clean_comment_empty():
    assert clean_comment("") == ""


# ---------- collect_merge_origins ----------


def test_collect_merge_origins():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.merge_cells("A1:B2")
    ws.merge_cells("D5:F7")
    origins = collect_merge_origins(ws)
    assert origins == {(1, 1), (5, 4)}
    wb.close()


# ---------- extract_single_sheet ----------


def test_extract_three_cells():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "Name"
    ws["B1"] = "Alice"
    ws["A2"] = "Age"
    cells = extract_single_sheet(ws, "Sheet1", "visible")
    assert len(cells) == 3
    # Check first cell
    assert cells[0][0] == 1  # row
    assert cells[0][1] == 1  # col
    assert cells[0][2] == "Name"  # value
    assert cells[0][5] == "Sheet1"  # sheet_name
    wb.close()


def test_cell_with_comment():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "Password"
    ws["A1"].comment = openpyxl.comments.Comment("secret note", "author")
    cells = extract_single_sheet(ws, "Sheet1", "visible")
    assert len(cells) == 1
    assert cells[0][4] == "secret note"  # comment field
    wb.close()


def test_cell_with_formula():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "=SUM(A1:A5)"
    cells = extract_single_sheet(ws, "Sheet1", "visible")
    assert len(cells) == 1
    assert cells[0][2] == "=SUM(A1:A5)"  # value
    assert cells[0][3] == "=SUM(A1:A5)"  # formula
    wb.close()


def test_hidden_sheet():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "data"
    cells = extract_single_sheet(ws, "SecretSheet", "hidden")
    assert cells[0][5] == "SecretSheet[HIDDEN]"
    wb.close()


def test_very_hidden_sheet():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "data"
    cells = extract_single_sheet(ws, "VerySecret", "veryHidden")
    assert cells[0][5] == "VerySecret[HIDDEN]"
    wb.close()


def test_merged_cell_origin():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "Merged"
    ws["B1"] = "NotMerged"
    ws.merge_cells("A1:A3")
    cells = extract_single_sheet(ws, "Sheet1", "visible")
    merged_cell = [c for c in cells if c[0] == 1 and c[1] == 1][0]
    not_merged = [c for c in cells if c[0] == 1 and c[1] == 2][0]
    assert merged_cell[6] is True  # is_merged_origin
    assert not_merged[6] is False
    wb.close()


def test_empty_cells_skipped():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "data"
    # B1, C1 are empty (None)
    ws["D1"] = "more"
    cells = extract_single_sheet(ws, "Sheet1", "visible")
    assert len(cells) == 2
    wb.close()


def test_comment_prefix_stripped():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "value"
    ws["A1"].comment = openpyxl.comments.Comment("Comment:\nactual text", "author")
    cells = extract_single_sheet(ws, "Sheet1", "visible")
    assert cells[0][4] == "actual text"
    wb.close()


# ---------- extract_workbook ----------


def test_extract_workbook_one_sheet(tmp_xlsx):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "Password"
    ws["B1"] = "s3cret!!"
    ws["A2"] = "Token"
    p = tmp_xlsx(wb)
    wb.close()

    sheets = extract_workbook(p)
    assert len(sheets) == 1
    assert len(sheets[0]) == 3


def test_extract_workbook_multiple_sheets(tmp_xlsx):
    wb = openpyxl.Workbook()
    ws1 = wb.active
    ws1.title = "Sheet1"
    ws1["A1"] = "data1"

    ws2 = wb.create_sheet("Sheet2")
    ws2["A1"] = "data2"
    ws2["B1"] = "data3"

    p = tmp_xlsx(wb)
    wb.close()

    sheets = extract_workbook(p)
    assert len(sheets) == 2
    assert len(sheets[0]) == 1
    assert len(sheets[1]) == 2


def test_extract_workbook_hidden_sheet(tmp_xlsx):
    wb = openpyxl.Workbook()
    ws1 = wb.active
    ws1.title = "Visible"
    ws1["A1"] = "open"

    ws2 = wb.create_sheet("Hidden")
    ws2["A1"] = "secret"
    ws2.sheet_state = "hidden"

    p = tmp_xlsx(wb)
    wb.close()

    sheets = extract_workbook(p)
    assert len(sheets) == 2
    # Hidden sheet should have [HIDDEN] suffix
    hidden_cells = sheets[1]
    assert hidden_cells[0][5] == "Hidden[HIDDEN]"


def test_extract_workbook_accepts_path_object(tmp_xlsx):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "test"
    p = tmp_xlsx(wb)
    wb.close()

    # Should accept both str and Path
    sheets_str = extract_workbook(str(p))
    sheets_path = extract_workbook(p)
    assert len(sheets_str) == len(sheets_path)


def test_extract_workbook_skips_chartsheet(tmp_xlsx):
    wb = openpyxl.Workbook()
    wb.active["A1"] = "Password"
    wb.active["B1"] = "s3cRet!99"
    chart = BarChart()
    chart.add_data(Reference(wb.active, min_col=2, min_row=1, max_row=1))
    wb.create_chartsheet("Chart1").add_chart(chart)
    p = tmp_xlsx(wb)
    wb.close()

    sheets = extract_workbook(p)
    assert len(sheets) == 1
    assert [c[2] for c in sheets[0]] == ["Password", "s3cRet!99"]


REL_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"


def _credential_workbook(tmp_xlsx) -> Path:
    wb = openpyxl.Workbook()
    wb.active["A1"] = "Password"
    wb.active["B1"] = "s3cRet!99"
    path = tmp_xlsx(wb)
    wb.close()
    return path


def _rewrite(src: Path, name: str, edits: dict[str, object]) -> Path:
    """Copy ``src`` to ``name`` beside it. ``edits`` maps a part to new bytes, or
    to a ``(old, new)`` replacement applied once to that part's XML."""
    dst = src.with_name(name)
    with zipfile.ZipFile(src) as zin, zipfile.ZipFile(dst, "w") as zout:
        present = set(zin.namelist())
        for item in zin.infolist():
            data = zin.read(item.filename)
            edit = edits.get(item.filename)
            if isinstance(edit, tuple):
                old, new = edit
                assert old in data, (item.filename, old)
                data = data.replace(old, new, 1)
            elif isinstance(edit, bytes):
                data = edit
            zout.writestr(item.filename, data)
        for part, data in edits.items():
            if part not in present:
                assert isinstance(data, bytes), part
                zout.writestr(part, data)
    return dst


def _assert_only_doormat_reads(path: Path) -> None:
    """openpyxl's own loader refuses the file; doormat still finds the credential."""
    with pytest.raises((KeyError, TypeError, ValueError)):
        openpyxl.load_workbook(path, data_only=False)
    sheets = extract_workbook(path)
    assert [c[2] for c in sheets[0]] == ["Password", "s3cRet!99"]


def test_extract_workbook_ignores_broken_external_link(tmp_xlsx):
    # The link part exists but its relationships part is missing.
    src = _credential_workbook(tmp_xlsx)
    path = _rewrite(src, "links.xlsx", {
        "xl/_rels/workbook.xml.rels": (
            b"</Relationships>",
            (
                b'<Relationship Id="rIdLink" Target="externalLinks/externalLink1.xml" '
                b'Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/'
                b'externalLink"/></Relationships>'
            ),
        ),
        "xl/workbook.xml": (
            b"</sheets>",
            f'</sheets><externalReferences><externalReference xmlns:r="{REL_NS}" '
            'r:id="rIdLink"/></externalReferences>'.encode(),
        ),
        "xl/externalLinks/externalLink1.xml": (
            '<externalLink xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" '
            f'xmlns:r="{REL_NS}"><externalBook r:id="rId1"/></externalLink>'.encode()
        ),
    })
    _assert_only_doormat_reads(path)


def test_extract_workbook_ignores_invalid_print_titles(tmp_xlsx):
    src = _credential_workbook(tmp_xlsx)
    path = _rewrite(src, "titles.xlsx", {
        "xl/workbook.xml": (
            b"<definedNames />",
            (
                b'<definedNames><definedName name="_xlnm.Print_Titles" '
                b'localSheetId="0">#N/A</definedName></definedNames>'
            ),
        ),
    })
    _assert_only_doormat_reads(path)


def test_extract_workbook_ignores_malformed_pivot_cache(tmp_xlsx):
    # A pivot cache with "-" where a number belongs, plus a pivot table using it.
    src = _credential_workbook(tmp_xlsx)
    main = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    pivot_rel = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/pivotTable"
    cache_rel = (
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/pivotCacheDefinition"
    )
    path = _rewrite(src, "pivot.xlsx", {
        "xl/_rels/workbook.xml.rels": (
            b"</Relationships>",
            f'<Relationship Id="rIdCache" Target="pivotCache/pivotCacheDefinition1.xml" '
            f'Type="{cache_rel}"/></Relationships>'.encode(),
        ),
        "xl/workbook.xml": (
            b"</workbook>",
            f'<pivotCaches><pivotCache xmlns:r="{REL_NS}" cacheId="1" r:id="rIdCache"/>'
            "</pivotCaches></workbook>".encode(),
        ),
        "xl/pivotCache/pivotCacheDefinition1.xml": (
            f'<pivotCacheDefinition xmlns="{main}"><cacheSource type="worksheet"/>'
            '<cacheFields count="1"><cacheField name="Amount"><sharedItems>'
            '<n v="-"/></sharedItems></cacheField></cacheFields></pivotCacheDefinition>'.encode()
        ),
        "xl/worksheets/_rels/sheet1.xml.rels": (
            f'<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
            f'<Relationship Id="rId1" Target="../pivotTables/pivotTable1.xml" Type="{pivot_rel}"/>'
            "</Relationships>".encode()
        ),
        "xl/pivotTables/pivotTable1.xml": (
            f'<pivotTableDefinition xmlns="{main}" name="Pivot1" cacheId="1" dataCaption="Values">'
            '<location ref="D1:E2" firstHeaderRow="1" firstDataRow="1" firstDataCol="1"/>'
            "</pivotTableDefinition>".encode()
        ),
    })
    _assert_only_doormat_reads(path)


def test_text_starting_with_equals_is_a_value_not_a_formula(tmp_xlsx):
    wb = openpyxl.Workbook()
    wb.active["A1"] = "Password"
    wb.active["B1"] = "=k9Xq!fake"
    wb.active["B1"].data_type = "s"
    wb.active["C1"] = "=SUM(1,2)"
    p = tmp_xlsx(wb)
    wb.close()

    cells = {c[1]: c[2:4] for c in extract_workbook(p)[0]}
    assert cells[2] == ("=k9Xq!fake", "")
    assert cells[3] == ("=SUM(1,2)", "=SUM(1,2)")


def test_extract_array_formula_as_text(tmp_xlsx):
    from openpyxl.worksheet.formula import ArrayFormula

    wb = openpyxl.Workbook()
    wb.active["A1"] = ArrayFormula("A1", "=SUM(B1:B2*C1:C2)")
    p = tmp_xlsx(wb)
    wb.close()

    cells = extract_workbook(p)[0]
    assert [c[2:4] for c in cells] == [("=SUM(B1:B2*C1:C2)", "=SUM(B1:B2*C1:C2)")]


def test_xml_entity_declarations_are_rejected(tmp_xlsx):
    """A worksheet that declares XML entities must be refused, not expanded."""
    import doormat

    wb = openpyxl.Workbook()
    wb.active["A1"] = "Password"
    wb.active["B1"] = "&e;"
    clean = tmp_xlsx(wb)
    wb.close()

    evil = clean.with_name("entities.xlsx")
    with zipfile.ZipFile(clean) as src, zipfile.ZipFile(evil, "w") as dst:
        for item in src.infolist():
            data = src.read(item.filename)
            if item.filename == "xl/worksheets/sheet1.xml":
                data = data.replace(b"<worksheet", b'<!DOCTYPE worksheet [<!ENTITY e "x">]><worksheet', 1)
                assert b"<!DOCTYPE" in data
            dst.writestr(item, data)

    doc = doormat.load(evil)
    assert [(e.sheet, e.kind) for e in doc.errors] == [("Sheet", "EntitiesForbidden")]
    assert doc.credentials() == []
    assert not doc.complete


def test_xml_entity_declarations_in_workbook_structure_raise(tmp_xlsx):
    import doormat

    wb = openpyxl.Workbook()
    wb.active["A1"] = "x"
    clean = tmp_xlsx(wb)
    wb.close()

    evil = _rewrite(clean, "entities_wb.xlsx", {
        "xl/workbook.xml": (b"<workbook", b'<!DOCTYPE workbook [<!ENTITY e "x">]><workbook'),
    })
    with pytest.raises(ValueError):
        doormat.load(evil)
