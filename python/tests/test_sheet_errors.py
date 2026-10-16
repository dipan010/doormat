"""A sheet that cannot be parsed is reported and the other sheets are still scanned."""

from __future__ import annotations

import zipfile
from pathlib import Path

import doormat
import openpyxl
import pytest


def _two_sheet_workbook(tmp_path: Path) -> Path:
    wb = openpyxl.Workbook()
    wb.active.title = "Broken"
    wb.active["A1"] = "x"
    good = wb.create_sheet("Good")
    good["A1"] = "Password"
    good["B1"] = "s3cRet!99"
    path = tmp_path / "two.xlsx"
    wb.save(path)
    wb.close()
    return path


def _break_part(src: Path, part: str, data: bytes) -> Path:
    dst = src.with_name("broken.xlsx")
    with zipfile.ZipFile(src) as zin, zipfile.ZipFile(dst, "w") as zout:
        for item in zin.infolist():
            zout.writestr(item, data if item.filename == part else zin.read(item.filename))
    return dst


def test_malformed_xlsx_sheet_is_reported_and_others_scanned(tmp_path: Path) -> None:
    path = _break_part(_two_sheet_workbook(tmp_path), "xl/worksheets/sheet1.xml", b"<worksheet")
    doc = doormat.load(path)

    assert [c.coordinate for c in doc.credentials(120)] == ["B1"]
    assert [(e.sheet, e.kind) for e in doc.errors] == [("Broken", "ParseError")]
    assert not doc.complete
    assert doc.sheet_count == 1


def test_complete_when_every_sheet_is_read(tmp_path: Path) -> None:
    doc = doormat.load(_two_sheet_workbook(tmp_path))
    assert doc.complete
    assert doc.errors == ()
    assert doc.sheet_count == 2


def test_sheet_error_repr_omits_message() -> None:
    err = doormat.SheetError(sheet="S", kind="ValueError", message="could not convert 'hunter2'")
    assert "hunter2" not in repr(err)
    assert repr(err) == "SheetError(sheet='S', kind='ValueError')"


def test_extractor_without_handler_raises(tmp_path: Path) -> None:
    from doormat.extract import extract_xlsx

    path = _break_part(_two_sheet_workbook(tmp_path), "xl/worksheets/sheet1.xml", b"<worksheet")
    with pytest.raises(SyntaxError):
        extract_xlsx(path)


def test_undecodable_csv_is_reported(tmp_path: Path) -> None:
    path = tmp_path / "latin.csv"
    path.write_bytes(b"Password,caf\xe9\n")
    doc = doormat.load(path)
    assert [(e.sheet, e.kind) for e in doc.errors] == [("latin", "ValueError")]
    assert doc.sheet_count == 0
    assert doc.credentials() == []


def test_malformed_ods_table_is_reported_and_others_scanned(tmp_path: Path) -> None:
    pytest.importorskip("odf")
    from odf.opendocument import OpenDocumentSpreadsheet
    from odf.table import Table, TableCell, TableRow
    from odf.text import P

    doc = OpenDocumentSpreadsheet()
    broken = Table(name="Broken")
    row = TableRow()
    cell = TableCell()
    cell.setAttrNS(
        "urn:oasis:names:tc:opendocument:xmlns:table:1.0", "number-columns-repeated", "many"
    )
    cell.addElement(P(text="x"))
    row.addElement(cell)
    broken.addElement(row)
    doc.spreadsheet.addElement(broken)

    good = Table(name="Good")
    row = TableRow()
    for text in ("Password", "s3cRet!99"):
        c = TableCell()
        c.addElement(P(text=text))
        row.addElement(c)
    good.addElement(row)
    doc.spreadsheet.addElement(good)
    path = tmp_path / "two.ods"
    doc.save(str(path))

    result = doormat.load(path)
    assert [(c.sheet, c.coordinate) for c in result.credentials(120)] == [("Good", "B1")]
    assert [(e.sheet, e.kind) for e in result.errors] == [("Broken", "ValueError")]


def test_failing_xls_sheet_is_reported_and_others_scanned(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    pytest.importorskip("xlrd")
    xlwt = pytest.importorskip("xlwt")
    from doormat.extract import xls

    wb = xlwt.Workbook()
    wb.add_sheet("Broken").write(0, 0, "x")
    good = wb.add_sheet("Good")
    good.write(0, 0, "Password")
    good.write(0, 1, "s3cRet!99")
    path = tmp_path / "two.xls"
    wb.save(str(path))

    real = xls._extract_sheet

    def fail_on_broken(sheet, budget=None):  # type: ignore[no-untyped-def]
        if sheet.name == "Broken":
            raise ValueError("bad record")
        return real(sheet, budget)

    monkeypatch.setattr(xls, "_extract_sheet", fail_on_broken)
    doc = doormat.load(path)
    assert [(c.sheet, c.coordinate) for c in doc.credentials(120)] == [("Good", "B1")]
    assert [(e.sheet, e.kind) for e in doc.errors] == [("Broken", "ValueError")]
